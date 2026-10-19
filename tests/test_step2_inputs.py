#!/usr/bin/env python3
"""
Step 2 input handling: validation, reading, and where every sample ends up.

    python3 tests/test_step2_inputs.py

Synthetic VCFs only; no reference set, no RAxML.  Covers the 3.37 changes:

  - a header-only VCF and a zero-coverage-only VCF are valid samples
    (identical to the reference), not "empty" or "unparseable"
  - a sample whose calls are all AC=1 is analysed and reported, not dropped
  - non-ASCII bytes in a header are tolerated and the file is left untouched
  - the reference check is deterministic and reads ##contig lines
  - a record with more columns than the header is read with the extra
    fields ignored, wherever it is in the file; a garbage line is dropped
    and counted, not parsed as position 0
  - blank lines are not reported as skipped records
  - a multi-sample VCF is refused unless --first_sample_only
  - a reference type with no metadata file runs
  - a crash inside the HTML-tree child process is an error, not a hang
  - the usability verdicts are not written to the stats workbook

and the corrections made after the 3.37 validation (report section 7):

  - D-02 column A of the workbooks keeps the 3.36 index format
  - D-03 an extra field on the first data line no longer rejects the file
  - D-04 an even split between two references aborts
  - D-05 trailing blank lines are not "records skipped"
  - D-10 an HTML tree that fails, by exception or by returning 1, is
    recorded in html_tree_failures
  - D-11 sample names such as 'NA', 'None' or 'null' stay names in the
    HTML tree
  - D-12 the step 1 no-SNP notice uses the step 2 SNP definition

and the corrections made after the 3.37 accuracy study (record A2, section 7):

  - L-2 a SNP freebayes reports inside a multi-base record is read from that
    record through its CIGAR, not shown as the reference base
  - L-3 a sample with passing SNPs whose confident mixed calls outnumber them
    more than 5 to 1 is flagged as a possible mixture, highlighted in the
    summary's Groupings table
  - D-06 a VCF that ends part-way through a record is excluded as truncated

and the long-read (nanopore) path (vsnp3_longread.py):

  - the evidence QUAL grows with supporting reads, falls with reference reads,
    scores a mixed site as confident, a thin site below the N band, and is on
    the freebayes scale (a clean site scores the per-read weight less the prior)
  - a written record's QUAL is never below the N band: a site the caller
    reported is N at the least, never the reference base
  - the genotype comes from the evidence model: a clonal site with a few
    reference-looking reads is homozygous whatever bcftools said, a half-and-half
    site is mixed, and the caller's genotype is kept in INFO when it differed
  - a one-strand call and a call beside a homopolymer run are limited to the N
    band, with the reason in FILTER and INFO
  - the rewritten VCF carries the platform and model in its header, keeps the
    caller's QUAL, drops indels and sub-threshold records, and is read by step 2,
    the step 1 counters and the group reporter exactly as a freebayes VCF is
"""

import hashlib
import os
import sys
import tempfile
import time

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BIN = os.path.join(REPO, 'bin')
sys.path.insert(0, BIN)

PASS, FAIL = [], []


def check(name, ok, detail=''):
    (PASS if ok else FAIL).append(name)
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f'\n        {detail}' if detail and not ok else ''))
    return ok


HEADER = (
    '##fileformat=VCFv4.2\n'
    '##reference=/data/ref/NC_1-NC_2.fasta\n'
    '##contig=<ID=NC_1,length=5000>\n'
    '##contig=<ID=NC_2,length=3000>\n'
    '##INFO=<ID=AC,Number=A,Type=Integer,Description="alt allele count">\n'
    '##INFO=<ID=DP,Number=1,Type=Integer,Description="depth">\n'
    '##INFO=<ID=MQM,Number=A,Type=Float,Description="mean mapping quality">\n'
    '#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tsample1\n'
)


def snp(chrom, pos, ref, alt, qual=900, ac=2, mq=60, dp=50):
    gt = '1/1' if ac == 2 else '0/1'
    return (f'{chrom}\t{pos}\t.\t{ref}\t{alt}\t{qual}\t.\t'
            f'AB=0;AC={ac};AN=2;DP={dp};MQM={mq};TYPE=snp\tGT:DP\t{gt}:{dp}\n')


def zc(chrom, pos):
    return f'{chrom}\t{pos}\t.\tN\t.\t.\t.\t.\tGT\t./.\n'


def write(directory, name, content):
    path = os.path.join(directory, name)
    mode = 'wb' if isinstance(content, bytes) else 'w'
    with open(path, mode) as f:
        f.write(content)
    return path


def sha(path):
    with open(path, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def make_set(directory):
    normal = HEADER + snp('NC_1', 100, 'A', 'G') + snp('NC_1', 200, 'C', 'T') + zc('NC_1', 300) + snp('NC_2', 50, 'G', 'A')
    files = {
        'normal_zc.vcf': normal,
        'allhet_zc.vcf': HEADER + snp('NC_1', 100, 'A', 'G', ac=1) + snp('NC_1', 200, 'C', 'T', ac=1) + zc('NC_1', 300),
        'headeronly_zc.vcf': HEADER,
        'zconly_zc.vcf': HEADER + zc('NC_1', 10) + zc('NC_1', 11) + zc('NC_2', 7),
        'nonascii_zc.vcf': normal.encode('ascii').replace(
            b'##reference=/data/ref/', '##reference=/data/café–lab/'.encode('utf-8') + b'\xe9'),
        'ragged_zc.vcf': HEADER + snp('NC_1', 100, 'A', 'G') + '\t'.join(['X'] * 20) + '\n' + snp('NC_1', 200, 'C', 'T'),
        'garbage_zc.vcf': os.urandom(2048),
        'otherref_zc.vcf': normal.replace('ID=NC_1,', 'ID=NZ_9,').replace('ID=NC_2,', 'ID=NZ_8,').replace('NC_1\t', 'NZ_9\t').replace('NC_2\t', 'NZ_8\t'),
        'multisample_zc.vcf': normal.replace('\tsample1\n', '\tsample1\tsample2\n').replace(':50\n', ':50\t0/0:40\n'),
        'zero_bytes_zc.vcf': '',
        # D-03: the same records as normal, one carrying an extra field.  On the
        # first data line pandas used to reject the whole file.
        'extrafirst_zc.vcf': (HEADER + snp('NC_1', 100, 'A', 'G').replace('\n', '\textra\n')
                              + snp('NC_1', 200, 'C', 'T') + zc('NC_1', 300) + snp('NC_2', 50, 'G', 'A')),
        'extralater_zc.vcf': (HEADER + snp('NC_1', 100, 'A', 'G') + snp('NC_1', 200, 'C', 'T')
                              + zc('NC_1', 300) + snp('NC_2', 50, 'G', 'A').replace('\n', '\textra\tmore\n')),
        # D-05
        'trailingblank_zc.vcf': normal + '\n\n',
    }
    return {name: write(directory, name, content) for name, content in files.items()}


def section_validator(directory, paths):
    print('\n[1] validation')
    from vsnp3_input_validator import InputValidator
    before = {name: sha(path) for name, path in paths.items()}
    vcf_list = sorted(paths.values())
    r1 = InputValidator().validate_vcf_list(vcf_list)
    r2 = InputValidator().validate_vcf_list(vcf_list)
    valid = {os.path.basename(p) for p in r1.valid_files}
    check('valid set is what it should be',
          valid == {'normal_zc.vcf', 'allhet_zc.vcf', 'headeronly_zc.vcf', 'zconly_zc.vcf',
                    'nonascii_zc.vcf', 'ragged_zc.vcf', 'multisample_zc.vcf',
                    'extrafirst_zc.vcf', 'extralater_zc.vcf', 'trailingblank_zc.vcf'},
          f'valid = {sorted(valid)}')
    check('header-only file is valid and listed as such',
          any(p.endswith('headeronly_zc.vcf') for p in r1.header_only_files))
    check('garbage is corrupted', any(p.endswith('garbage_zc.vcf') for p in r1.corrupted_files),
          f'{r1.corrupted_files}')
    check('zero-byte file is empty', any(p.endswith('zero_bytes_zc.vcf') for p in r1.empty_files))
    check('other reference is a mismatch, not a crash',
          [os.path.basename(m['file']) for m in r1.reference_mismatches] == ['otherref_zc.vcf'],
          f'{r1.reference_mismatches}')
    check('reference comes from ##contig lines', r1.reference_contigs == ('NC_1', 'NC_2'),
          f'{r1.reference_contigs}')
    check('a minority on another reference does not lose the majority',
          not r1.reference_majority_lost)
    check('non-ASCII header bytes are noted, not an error',
          [os.path.basename(p) for p in r1.non_ascii_files] == ['nonascii_zc.vcf'],
          f'{r1.non_ascii_files}')
    check('every excluded file has a reason',
          all(r1.reason(p) for p in vcf_list if p not in r1.valid_files))
    check('validation is deterministic',
          (r1.valid_files, r1.reasons, r1.reference_contigs) == (r2.valid_files, r2.reasons, r2.reference_contigs))
    after = {name: sha(path) for name, path in paths.items()}
    check('validation modifies no input file', before == after)


def section_reader(directory, paths):
    print('\n[2] reading')
    from vsnp3_step2 import parse_vcf
    before = {name: sha(path) for name, path in paths.items()}

    _, df, stats, err = parse_vcf(paths['normal_zc.vcf'])
    check('normal VCF parses', err is None and df is not None, str(err))
    if df is not None:
        check('columns are what step 2 expects',
              list(df.columns) == ['CHROM', 'POS', 'REF', 'ALT', 'QUAL', 'AC', 'DP', 'MQ', 'CIGAR', 'single_base', 'abs_pos'],
              str(list(df.columns)))
        row = df[df['abs_pos'] == 'NC_1:100'].iloc[0]
        check('AC, DP and MQ (from MQM) are read', (row['AC'], row['DP'], row['MQ']) == (2, 50, 60),
              f'{row.to_dict()}')
        zrow = df[df['abs_pos'] == 'NC_1:300'].iloc[0]
        check('zero-coverage record keeps QUAL and MQ as missing, not zero',
              bool(zrow[['QUAL', 'MQ']].isna().all()) and zrow['AC'] == 0)
        check('stats count called, zero-coverage and AC',
              (stats['records'], stats['called'], stats['zero_coverage'], stats['ac2'], stats['ac1']) == (4, 3, 1, 3, 0),
              str(stats))

    _, df, stats, err = parse_vcf(paths['allhet_zc.vcf'])
    check('all-AC=1 VCF parses', err is None and stats['ac1'] == 2 and stats['ac2'] == 0, f'{err} {stats}')

    _, df, stats, err = parse_vcf(paths['headeronly_zc.vcf'])
    check('header-only VCF parses to an empty frame with the right columns',
          err is None and len(df) == 0 and 'abs_pos' in df.columns, str(err))

    _, df, stats, err = parse_vcf(paths['zconly_zc.vcf'])
    check('zero-coverage-only VCF parses (no MQ guard false alarm)',
          err is None and len(df) == 3 and stats['called'] == 0, f'{err} {stats}')

    _, df, stats, err = parse_vcf(paths['nonascii_zc.vcf'])
    check('non-ASCII bytes are dropped on read and counted',
          err is None and stats['non_ascii_bytes'] == 6 and len(df) == 4, f'{err} {stats}')

    # Twenty fields of 'X': too many columns AND no numeric POS.  It is the POS
    # that gets it dropped and counted; the extra columns alone would not.
    _, df, stats, err = parse_vcf(paths['ragged_zc.vcf'])
    check('a garbage line (no numeric POS, too many fields) is dropped and counted, rest kept',
          err is None and len(df) == 2 and stats['skipped_lines'] + stats['malformed'] == 1,
          f'{err} {stats} {None if df is None else df["abs_pos"].tolist()}')
    if df is not None:
        check('no position-0 record from the skipped line', not (df['POS'] == 0).any())

    # D-03.  An extra field is ignored wherever the record is: same frame as
    # the file without it, and nothing reported as skipped.
    import pandas as pd
    _, ref_df, _, _ = parse_vcf(paths['normal_zc.vcf'])
    for name, where in (('extrafirst_zc.vcf', 'the first data line'),
                        ('extralater_zc.vcf', 'a later line')):
        _, df, stats, err = parse_vcf(paths[name])
        same = False
        if df is not None:
            try:
                pd.testing.assert_frame_equal(df.reset_index(drop=True),
                                              ref_df.reset_index(drop=True), check_exact=True)
                same = True
            except AssertionError as e:
                err = f'{err} {e}'
        check(f'an extra field on {where} is ignored: same records, nothing counted (D-03)',
              err is None and same and stats['skipped_lines'] == 0 and stats['malformed'] == 0,
              f'{err} {stats}')

    # D-05.  Blank lines are not records, so they are not "skipped records".
    _, df, stats, err = parse_vcf(paths['trailingblank_zc.vcf'])
    check('trailing blank lines are not reported as skipped records (D-05)',
          err is None and len(df) == 4 and stats['skipped_lines'] == 0 and stats['malformed'] == 0,
          f'{err} {stats}')
    from vsnp3_step2 import dataframe_from_vcf_body
    # The body as normalize_vcf_text hands it over: MQM already renamed MQ.
    record = snp('NC_1', 100, 'A', 'G').replace(';MQM=', ';MQ=')
    other = snp('NC_1', 200, 'C', 'T').replace(';MQM=', ';MQ=')
    blanks = {'leading': '\n' + record, 'between': record + '\n\n' + other,
              'several trailing': record + '\n\n\n\n'}
    counts = {k: dataframe_from_vcf_body([], body)[1]['skipped_lines'] for k, body in blanks.items()}
    check('blank lines anywhere in the body are not counted (D-05)',
          set(counts.values()) == {0}, str(counts))

    _, df, stats, err = parse_vcf(paths['multisample_zc.vcf'])
    check('multi-sample VCF is refused with a reason', df is None and 'multi-sample' in str(err), str(err))
    _, df, stats, err = parse_vcf(paths['multisample_zc.vcf'], first_sample_only=True)
    check('--first_sample_only accepts it', err is None and stats['multi_sample'] == 2, f'{err} {stats}')

    _, df, stats, err = parse_vcf(paths['garbage_zc.vcf'])
    check('garbage is an error, not an exception', df is None and err)

    after = {name: sha(path) for name, path in paths.items()}
    check('reading modifies no input file', before == after)


def section_fix_vcfs(directory, paths):
    print('\n[3] normalized copies on request only')
    from vsnp3_step2 import parse_vcf
    out = os.path.join(directory, 'normalized')
    os.makedirs(out, exist_ok=True)
    parse_vcf(paths['normal_zc.vcf'], normalized_dir=out)
    copy = os.path.join(out, 'normal_zc.vcf')
    ok = os.path.exists(copy)
    check('normalized copy written when asked', ok)
    if ok:
        body = open(copy).read()
        check('normalized copy renames MQM to MQ', ';MQ=60' in body and 'MQM' not in body.split('#CHROM')[1])


def section_group(directory, paths):
    print('\n[4] grouping keeps and explains no-SNP samples')
    import pandas as pd
    from vsnp3_group_on_defining_snps import Group
    from vsnp3_step2 import parse_vcf

    group = Group.__new__(Group)
    group.debug = False
    group.n_threshold, group.qual_threshold, group.mq_threshold = 50, 150, 56
    group.ambigious_lookup = {'AG': 'R', 'CT': 'Y', 'GA': 'R', 'TC': 'Y'}
    group.sample_stats = {}

    _, allhet, _, _ = parse_vcf(paths['allhet_zc.vcf'])
    _, normal, _, _ = parse_vcf(paths['normal_zc.vcf'])
    _, empty, _, _ = parse_vcf(paths['headeronly_zc.vcf'])
    _, zconly, _, _ = parse_vcf(paths['zconly_zc.vcf'])

    check('normal sample has passing SNPs', len(group._passing(normal)) == 3)
    check('all-AC=1 sample has none', len(group._passing(allhet)) == 0)
    d = group.describe_no_snps(allhet, 'allhet_zc.vcf')
    check('all-AC=1 sample is flagged as mixed', d['mixed_only'] and '2 mixed (AC=1)' in d['reason'], str(d))
    d = group.describe_no_snps(empty, 'headeronly_zc.vcf')
    check('header-only sample is explained as reference-identical',
          'identical to the reference' in d['reason'] and not d['mixed_only'], str(d))
    d = group.describe_no_snps(zconly, 'zconly_zc.vcf')
    check('zero-coverage-only sample is explained', 'all zero-coverage' in d['reason'], str(d))

    # A reference-identical sample sits in a group beside informative ones and
    # comes out as a row of reference calls, the same length as the others.
    group.dataframes_names_updated = {'normal': normal, 'allhet': allhet, 'empty': empty}
    essentials = {'normal': group._passing(normal), 'allhet': group._passing(allhet),
                  'empty': group._passing(empty)}
    # Two informative samples are needed for a position to be informative.
    other = normal.copy()
    other.loc[other['abs_pos'] == 'NC_1:100', 'ALT'] = 'A'   # differs from normal at 100
    group.dataframes_names_updated['other'] = other
    essentials['other'] = group._passing(other)
    result, name = group.make_groupings(('G', essentials))
    seqs = {s: ''.join(df['ALT'].tolist()) for s, df in result.items()}
    check('group with a no-SNP sample still aligns', set(seqs) == {'normal', 'allhet', 'empty', 'other', 'root'},
          str(seqs))
    if 'root' in seqs:
        check('reference-identical sample equals root', seqs.get('empty') == seqs['root'], str(seqs))
        check('all-AC=1 sample carries IUPAC codes, not reference calls',
              seqs.get('allhet') != seqs['root'] and set(seqs.get('allhet', '')) <= set('RYSWKMN-ACGT'), str(seqs))
        check('every row is the same length', len({len(s) for s in seqs.values()}) == 1, str(seqs))

    all_empty = {'empty': group._passing(empty), 'zconly': group._passing(zconly)}
    group.dataframes_names_updated.update({'zconly': zconly})
    result, name = group.make_groupings(('E', all_empty))
    check('a group of only no-SNP samples is skipped, not fatal', result == {})


def section_metadata():
    print('\n[5] reference type without a metadata file')
    from vsnp3_group_on_defining_snps import Group
    from vsnp3_reference_options import Ref_Options
    name, key, hits = Group.resolve_sample_name_detail('x_zc.vcf', None, False)
    check('name resolution with no metadata frame', (name, key, hits) == ('x_zc.vcf', None, 0))
    with tempfile.TemporaryDirectory() as tmp:
        # A path containing "meta" must not hide the defining-SNP file.
        ref_dir = os.path.join(tmp, 'metagenomics', 'Bug_ref')
        os.makedirs(ref_dir)
        for f in ('Bug_define_filter.xlsx', 'Bug.fasta', 'Bug.gbk'):
            open(os.path.join(ref_dir, f), 'w').close()
        ro = Ref_Options.__new__(Ref_Options)
        ro.select_ref = 'Bug_ref'
        ro.metadata_gather(ref_dir)
        check('defining-SNP file found under a "meta" directory',
              ro.defining_snps is not None and ro.defining_snps.endswith('Bug_define_filter.xlsx'),
              str(ro.defining_snps))
        check('no metadata file means metadata is None, not an error', ro.metadata is None)
        try:
            Group.read_metadata(os.path.join(ref_dir, 'missing_metadata.xlsx'))
            check('read_metadata raises on a missing file', False)
        except Exception:                                        # noqa: BLE001
            check('read_metadata raises on a missing file (caught by Group)', True)


def _boom():
    raise ValueError('child process failure')


def _slow():
    time.sleep(30)


def _fine():
    return 42


def section_timeout():
    print('\n[6] HTML tree child process outcomes')
    from vsnp3_group_on_defining_snps import run_with_timeout
    check('result comes back', run_with_timeout(_fine, 20) == 42)
    started = time.monotonic()
    try:
        run_with_timeout(_boom, 20)
        check('a crash in the child is an error, not a hang', False)
    except RuntimeError as e:
        check('a crash in the child is an error, not a hang',
              time.monotonic() - started < 15 and 'child process failure' in str(e), str(e))
    started = time.monotonic()
    try:
        run_with_timeout(_slow, 2)
        check('timeout still raises TimeoutError', False)
    except TimeoutError:
        check('timeout still raises TimeoutError', time.monotonic() - started < 15)


def _other_reference(text):
    '''The same VCF called against a reference with other contig names.'''
    return (text.replace('ID=NC_1,', 'ID=NZ_9,').replace('ID=NC_2,', 'ID=NZ_8,')
            .replace('NC_1\t', 'NZ_9\t').replace('NC_2\t', 'NZ_8\t'))


def _third_reference(text):
    return (text.replace('ID=NC_1,', 'ID=CP_5,').replace('ID=NC_2,', 'ID=CP_6,')
            .replace('NC_1\t', 'CP_5\t').replace('NC_2\t', 'CP_6\t'))


def section_reference_split(directory):
    print('\n[8] reference majority (D-04)')
    from vsnp3_input_validator import InputValidator
    from vsnp3_step2 import VCF_to_DF
    body = HEADER + snp('NC_1', 100, 'A', 'G') + snp('NC_2', 50, 'G', 'A')
    sets = {
        'even 2+2': [body, body, _other_reference(body), _other_reference(body)],
        'even 3+3': [body] * 3 + [_other_reference(body)] * 3,
        'three-way 2+2+2': [body] * 2 + [_other_reference(body)] * 2 + [_third_reference(body)] * 2,
        'majority 3+2': [body] * 3 + [_other_reference(body)] * 2,
    }
    results = {}
    for label, contents in sets.items():
        d = os.path.join(directory, 'split_' + label.replace(' ', '_').replace('+', '_'))
        os.makedirs(d)
        files = [write(d, f's{i}_zc.vcf', c) for i, c in enumerate(contents)]
        results[label] = (files, InputValidator().validate_vcf_list(sorted(files)))
    for label in ('even 2+2', 'even 3+3', 'three-way 2+2+2'):
        r = results[label][1]
        check(f'{label}: no strict majority, so the set is reported as mixed',
              r.reference_majority_lost and any('not a majority' in n for n in r.notes),
              f'lost={r.reference_majority_lost} notes={r.notes}')
    files, r = results['majority 3+2']
    check('3+2: the majority is kept and the 2 others excluded and named',
          not r.reference_majority_lost and len(r.valid_files) == 3
          and len(r.reference_mismatches) == 2, f'{r.valid_files} {r.reference_mismatches}')

    # The abort itself, as step 2 runs it: nothing analysed, exit status 1.
    files = results['even 3+3'][0]
    cwd = os.getcwd()
    try:
        os.chdir(os.path.dirname(files[0]))
        try:
            VCF_to_DF(vcf_list=sorted(files))
            check('step 2 stops on an even split', False, 'VCF_to_DF returned')
        except SystemExit as e:
            check('step 2 stops on an even split', e.code == 1, f'exit code {e.code}')
    finally:
        os.chdir(cwd)


def _table_frame(names):
    import pandas as pd
    rows = len(names) - 3                               # root, samples..., MQ, annotation
    bases = {'NC_1:100': 'AGGAAGAGAG', 'NC_1:200': 'CTTCTCTCTC',
             'NC_1:300': 'GGAAAGAGGA', 'NC_1:400': 'TTTCCTCTCT'}
    data = {pos: list(seq[:rows + 1]) + [60, f'gene {pos}'] for pos, seq in bases.items()}
    return pd.DataFrame(data, index=names)


def _tables():
    from vsnp3_fasta_to_snps_table import Tables
    tables = Tables.__new__(Tables)
    tables.show_groups = False
    tables.st = 'test'
    return tables


def section_workbook_format(directory):
    print('\n[9] workbook column A format (D-02)')
    try:
        import openpyxl
    except ImportError as e:
        print(f'  SKIP  ({e})')
        return
    names = ['root', 'S1', 'S2', 'MQ', 'annotation']
    tables = _tables()
    path = os.path.join(directory, 'format_table.xlsx')
    tables.excel_formatter(tables.json_round_trip(_table_frame(names)), path)
    ws = openpyxl.load_workbook(path).active
    col_a = [ws.cell(row, 1) for row in range(2, 2 + len(names))]
    check('column A holds the row labels', [c.value for c in col_a] == names,
          str([c.value for c in col_a]))

    def index_style(cell):
        b = cell.border
        return (bool(cell.font.b), b.left.style, b.right.style, b.top.style, b.bottom.style,
                cell.alignment.horizontal, cell.alignment.vertical)
    expected = (True, 'thin', 'thin', 'thin', 'thin', 'center', 'top')
    styles = [index_style(c) for c in col_a]
    check('column A is bold, bordered, centred and top-aligned, as 3.36 wrote it (D-02)',
          all(s == expected for s in styles), str(styles))
    label, cell = ws.cell(1 + len(names), 1), ws.cell(1 + len(names), 2)
    color = label.font.color.rgb if label.font.color is not None else None
    check('the annotation label does not take the annotation row format (D-02)',
          not label.alignment.textRotation and color != 'FF0A028C',
          f'rotation {label.alignment.textRotation}, colour {color}')
    check('the annotation cells keep the rotated blue row format',
          cell.alignment.textRotation == 180 and cell.font.color.rgb == 'FF0A028C')
    check('position cells are not given the column A format',
          not ws.cell(3, 2).font.b and ws.cell(3, 2).border.left.style is None)


def section_na_names(directory):
    print('\n[10] sample names read as missing by default (D-11)')
    try:
        import vsnp3_html_tree as vht
    except ImportError as e:
        print(f'  SKIP  ({e})')
        return
    names = ['root', 'NA', 'None', 'null', 'S4', 'MQ', 'annotation']
    tables = _tables()
    table = tables.json_round_trip(_table_frame(names))
    handed_over = tables.as_read_back(table)
    check('the frame step 2 hands to html_tree keeps NA-like names (D-11)',
          handed_over.iloc[:, 0].tolist() == names, str(handed_over.iloc[:, 0].tolist()))
    path = os.path.join(directory, 'na_names_table.xlsx')
    tables.excel_formatter(table, path)
    read_back = vht.read_table(path)
    check('read_table keeps them when html_tree reads a workbook (D-11)',
          read_back.iloc[:, 0].tolist() == names, str(read_back.iloc[:, 0].tolist()))
    import pandas as pd
    a, b = read_back.to_numpy(dtype=object), handed_over.to_numpy(dtype=object)
    same = a.shape == b.shape and bool(((a == b) | (pd.isna(a) & pd.isna(b))).all())
    check('read_table and as_read_back still agree', same and
          list(read_back.columns) == list(handed_over.columns))
    matrix = vht.SnpMatrix(vht.select_position_columns(handed_over))
    check('each such sample is a leaf under its own name, with its own clade bit',
          matrix.sample_names == ['NA', 'None', 'null', 'S4']
          and len(set(vht.leaf_bits(matrix.sample_names).values())) == 4,
          str(matrix.sample_names))


def _returns_one(*args, **kwargs):
    return 1


def section_html_tree_failures(directory):
    print('\n[11] HTML tree failures are recorded (D-10)')
    try:
        import pandas as pd
        import vsnp3_group_on_defining_snps as gods
        import vsnp3_html_tree  # noqa: F401  (plotly)
    except ImportError as e:
        print(f'  SKIP  ({e})')
        return

    class FakeTables:
        def __init__(self, path, df):
            self.table_to_tree_path, self.table_to_tree_df = path, df

    group = gods.Group.__new__(gods.Group)
    group.html_tree_failures = {}

    # A table html_tree cannot draw: no position columns.
    bad = pd.DataFrame({'Unnamed: 0': ['root', 'S1', 'S2'], 'x': ['A', 'A', 'C']})
    group.html_tree_build('G-exception', FakeTables(os.path.join(directory, 'bad.xlsx'), bad))
    reason = group.html_tree_failures.get('G-exception', '')
    check('an exception in html_tree is recorded with its message (D-10)',
          'NoPositionColumnsError' in reason, str(group.html_tree_failures))

    # html_tree as it behaves without raise_errors: prints, returns 1.
    real = gods.html_tree
    try:
        gods.html_tree = _returns_one
        group.html_tree_build('G-returned-1', FakeTables(os.path.join(directory, 'bad.xlsx'), bad))
    finally:
        gods.html_tree = real
    check('an html_tree that returns 1 is recorded (D-10)',
          'returned 1' in group.html_tree_failures.get('G-returned-1', ''),
          str(group.html_tree_failures))

    tables = _tables()
    names = ['root', 'NA', 'None', 'null', 'S4', 'MQ', 'annotation']
    table = tables.json_round_trip(_table_frame(names))
    path = os.path.join(directory, 'good_table.xlsx')
    tables.excel_formatter(table, path)
    group.html_tree_build('G-good', FakeTables(path, tables.as_read_back(table)))
    html_path = os.path.join(directory, 'good_table_position_tree.html')
    check('a tree that is drawn is not recorded, and the HTML is written',
          'G-good' not in group.html_tree_failures and os.path.exists(html_path),
          str(group.html_tree_failures))

    from vsnp3_html_tree import html_tree
    check('html_tree on the command line still prints and returns 1',
          html_tree(os.path.join(directory, 'bad.xlsx'), df=bad) == 1)


def section_step1_notice(directory):
    print('\n[12] step 1 no-SNP notice uses the step 2 definition (D-12)')
    import pandas as pd
    from vsnp3_zero_coverage import step2_snp_counts, no_snp_notice
    from vsnp3_group_on_defining_snps import Group
    from vsnp3_step2 import parse_vcf
    group = Group.__new__(Group)
    group.qual_threshold, group.mq_threshold = 150, 56

    # Boundaries on both sides of every threshold, as freebayes writes them.
    boundary = (HEADER
                + snp('NC_1', 10, 'A', 'G', qual=151, mq=56)        # passes
                + snp('NC_1', 11, 'A', 'G', qual=150, mq=60)        # QUAL not above 150
                + snp('NC_1', 12, 'A', 'G', qual='150.9', mq=60)    # truncated to 150
                + snp('NC_1', 13, 'A', 'G', qual=900, mq='55.9')    # MQ truncated to 55
                + snp('NC_1', 14, 'A', 'G', qual=900, mq='56.0')    # passes
                + snp('NC_1', 15, 'A', 'G', qual=900, ac=1)         # mixed
                + snp('NC_1', 16, 'AT', 'GC', qual=900)             # not single-base
                + snp('NC_1', 17, 'A', 'G', qual=900).replace(';MQM=60', '')  # no MQ
                + zc('NC_1', 18))
    vcfs = {'boundary': write(directory, 'boundary_zc.vcf', boundary)}
    for name in ('normal', 'allhet', 'zconly', 'headeronly'):
        vcfs[name] = os.path.join(directory, f'{name}_zc.vcf')
    agree = {}
    for name, path in vcfs.items():
        try:
            vcf_df = pd.read_csv(path, sep='\t', header=None, comment='#',
                                 names=['CHROM', 'POS', 'ID', 'REF', 'ALT', 'QUAL', 'FILTER',
                                        'INFO', 'FORMAT', 'Sample'])
        except pd.errors.EmptyDataError:
            vcf_df = pd.DataFrame(columns=['CHROM', 'POS', 'ID', 'REF', 'ALT', 'QUAL', 'FILTER',
                                           'INFO', 'FORMAT', 'Sample'])
        _, step2_df, _, _ = parse_vcf(path)
        agree[name] = (step2_snp_counts(vcf_df)[1], len(group._passing(step2_df)))
    check('step 1 counts exactly what step 2 takes as SNP positions',
          all(a == b for a, b in agree.values()) and agree['boundary'][0] == 2, str(agree))

    # CT14b: 20 high-quality AC=1 calls, no AC=2.  "Quality SNPs" was 20, so
    # the old notice (good_snp_count == 0) never appeared.
    mixed = no_snp_notice(step2_snp_count=0, ac1_count=20, ac2_snp_count=0)
    check('a mixed sample with high-quality AC=1 calls gets the mixed notice (D-12)',
          mixed is not None and 'mixed' in mixed and 'AC=2' in mixed, str(mixed))
    low = no_snp_notice(step2_snp_count=0, ac1_count=0, ac2_snp_count=5)
    check('AC=2 calls that all fail QUAL or MQ are not called reference-identical',
          low is not None and 'matches the reference' not in low, str(low))
    ref = no_snp_notice(step2_snp_count=0, ac1_count=0, ac2_snp_count=0)
    check('no call at all reads as matching the reference',
          ref is not None and 'matches the reference' in ref, str(ref))
    check('no notice when step 2 has SNPs, or the count is unknown',
          no_snp_notice(3, 20, 3) is None and no_snp_notice(None, 0, 0) is None)


def complex_rec(chrom, pos, ref, alt, qual=900, ac=2, cigar=None, mq=60):
    gt = '1/1' if ac == 2 else '0/1'
    cig = f'CIGAR={cigar};' if cigar else ''
    return (f'{chrom}\t{pos}\t.\t{ref}\t{alt}\t{qual}\t.\t'
            f'AB=0;AC={ac};AN=2;{cig}DP=50;MQM={mq};TYPE=complex\tGT:DP\t{gt}:50\n')


def section_multibase(directory):
    print('\n[13] a SNP inside a multi-base record is read from the record (L-2)')
    from vsnp3_group_on_defining_snps import Group
    from vsnp3_step2 import parse_vcf
    # Reference context NC_1:100-103 = C A C T.  "single" carries A>C at 101 as
    # its own record; every other sample reports 101 only inside a record that
    # starts at 100, which is how freebayes writes it next to an indel.
    cases = {
        'single':      (snp('NC_1', 101, 'A', 'C'), 'C'),
        'refsample':   (snp('NC_1', 500, 'T', 'C'), 'A'),
        'mnp':         (complex_rec('NC_1', 100, 'CAC', 'CCC', qual=162, cigar='1M1X1M'), 'C'),
        'mnp_het':     (complex_rec('NC_1', 100, 'CAC', 'CCC', ac=1, cigar='1M1X1M'), 'M'),
        'mnp_lowqual': (complex_rec('NC_1', 100, 'CAC', 'CCC', qual=40, cigar='1M1X1M'), 'A'),
        'mnp_midqual': (complex_rec('NC_1', 100, 'CAC', 'CCC', qual=100, cigar='1M1X1M'), 'N'),
        'deletion':    (complex_rec('NC_1', 100, 'CACT', 'CT', cigar='1M2D1M'), 'N'),
        'no_cigar':    (complex_rec('NC_1', 100, 'CAC', 'CCC'), 'C'),
        'no_cigar_indel': (complex_rec('NC_1', 100, 'CA', 'C'), 'N'),
        'two_alleles': (complex_rec('NC_1', 100, 'CAC', 'CCC,CTC', cigar='1M1X1M,1M1X1M'), 'N'),
        'change_at_start_only': (complex_rec('NC_1', 100, 'CAC', 'GAC', cigar='1X2M'), 'A'),
        'insertion':   (complex_rec('NC_1', 100, 'CA', 'CGGA', cigar='1M2I1M'), 'A'),
    }
    frames = {}
    for name, (record, _expected) in cases.items():
        path = write(directory, f'mb_{name}_zc.vcf', HEADER + record + snp('NC_1', 400, 'G', 'A'))
        _, frames[name], _, _ = parse_vcf(path)
    check('the reader keeps the CIGAR of multi-base records only',
          frames['mnp'].set_index('abs_pos').at['NC_1:100', 'CIGAR'] == '1M1X1M'
          and frames['single']['CIGAR'].isna().all(), str(frames['mnp']))
    group = Group.__new__(Group)
    group.debug = False
    group.n_threshold, group.qual_threshold, group.mq_threshold = 50, 150, 56
    group.ambigious_lookup = {'AG': 'R', 'GA': 'R', 'CT': 'Y', 'TC': 'Y', 'GC': 'S', 'CG': 'S',
                              'AT': 'W', 'TA': 'W', 'GT': 'K', 'TG': 'K', 'AC': 'M', 'CA': 'M'}
    group.dataframes_names_updated = frames
    essentials = {n: group._passing(df) for n, df in frames.items()}
    result, _ = group.make_groupings(('G', essentials))
    calls = {n: dict(zip(df['abs_pos'], df['ALT'])).get('NC_1:101') for n, df in result.items()}
    check('position 101 is in the table', all(v is not None for v in calls.values()), str(calls))
    check('the root is the reference base', calls.get('root') == 'A', str(calls))
    for name, (_record, expected) in cases.items():
        check(f'{name}: {expected}', calls.get(name) == expected, f'got {calls.get(name)!r}')
    # Eight cases differ from the reference at 101 inside a record; the
    # insertion and the change at the record's first base leave 101 as it is.
    check('every cell read from a record is counted',
          sum(group.embedded_cells.values()) == 8, str(dict(group.embedded_cells)))
    # Everyone carrying the same base makes the position uninformative; it then
    # leaves the table, and that is reported.
    group2 = Group.__new__(Group)
    group2.debug = False
    group2.n_threshold, group2.qual_threshold, group2.mq_threshold = 50, 150, 56
    group2.ambigious_lookup = group.ambigious_lookup
    same = {'single': frames['single'], 'mnp': frames['mnp'], 'no_cigar': frames['no_cigar']}
    group2.dataframes_names_updated = same
    same_ess = {n: group2._passing(df) for n, df in same.items()}
    result2, _ = group2.make_groupings(('H', same_ess))
    positions = set(result2['root']['abs_pos']) if result2 else set()
    check('a position all samples share once read correctly leaves the table, and is reported',
          'NC_1:101' not in positions and 'NC_1:101' in group2.embedded_now_uninformative['H'],
          f'{positions} {dict(group2.embedded_now_uninformative)}')


def section_mixed_signature(directory):
    print('\n[14] mixed-sample signature with passing SNPs (L-3)')
    import pandas as pd
    from vsnp3_version import mixed_signature, MIXED_MIN_AC1
    from vsnp3_zero_coverage import confident_mixed_count, mixed_sample_notice
    from vsnp3_group_on_defining_snps import Group
    from vsnp3_step2 import parse_vcf
    check('rule: at least the minimum, and more than 5 mixed calls per passing SNP',
          mixed_signature(0, MIXED_MIN_AC1) and not mixed_signature(0, MIXED_MIN_AC1 - 1)
          and mixed_signature(1, 87) and not mixed_signature(20, 100) and mixed_signature(20, 101)
          and not mixed_signature(2, MIXED_MIN_AC1 - 1) and not mixed_signature(460, 173))
    # The B. suis run that set the limits: the six clear ones are kept, the
    # seven at 1.0-2.8 calls per passing SNP are no longer flagged.
    kept = [(5, 58), (54, 467), (61, 338), (60, 591), (63, 721), (57, 775)]
    dropped = [(19, 44), (57, 58), (20, 39), (7, 11), (19, 54), (29, 30), (52, 53)]
    check('rule: the B. suis run keeps its 6 clear mixtures and drops the 7 marginal ones',
          all(mixed_signature(p, c) for p, c in kept)
          and not any(mixed_signature(p, c) for p, c in dropped))
    body = ''.join(snp('NC_1', 100 + i, 'A', 'G', ac=1) for i in range(12))
    body += snp('NC_1', 200, 'A', 'G', ac=1, qual=149) + snp('NC_1', 201, 'A', 'G', ac=1, mq=50)
    body += snp('NC_1', 300, 'C', 'T') + snp('NC_1', 301, 'C', 'T')
    path = write(directory, 'mixture_zc.vcf', HEADER + body)
    vcf_df = pd.read_csv(path, sep='\t', header=None, comment='#',
                         names=['CHROM', 'POS', 'ID', 'REF', 'ALT', 'QUAL', 'FILTER', 'INFO', 'FORMAT', 'Sample'])
    _, step2_df, _, _ = parse_vcf(path)
    group = Group.__new__(Group)
    group.qual_threshold, group.mq_threshold = 150, 56
    check('step 1 and step 2 count the same confident mixed calls',
          confident_mixed_count(vcf_df) == group._confident_mixed(step2_df) == 12,
          f'{confident_mixed_count(vcf_df)} {group._confident_mixed(step2_df)}')
    note = mixed_sample_notice(step2_snp_count=2, confident_ac1=25)
    check('a mixture with passing SNPs gets the step 1 notice',
          note is not None and 'Possible mixed sample' in note, str(note))
    check('no mixed notice for a clonal sample, a no-SNP sample (other notice), too few '
          'mixed calls or unknown counts',
          mixed_sample_notice(460, 173) is None and mixed_sample_notice(0, 50) is None
          and mixed_sample_notice(2, 12) is None
          and mixed_sample_notice(None, 50) is None and mixed_sample_notice(5, None) is None)


def section_summary_layout():
    print('\n[18] step 2 summary: mixed samples in the Groupings table, three sections gone')
    from types import SimpleNamespace
    from vsnp3_step2 import HTML_Summary, SampleAccounting
    summary = HTML_Summary.__new__(HTML_Summary)
    summary.out = []
    summary.group = SimpleNamespace(
        mixed_samples={'M1': {'file': 'M1.vcf', 'passing': 5, 'confident_ac1': 58}},
        qual_threshold=150, mq_threshold=56, group_failures={},
        embedded_cells={'C1': 3}, embedded_now_uninformative={'G1': {'NC_1:5'}})
    accounting = SampleAccounting(['M1.vcf', 'C1.vcf'])
    accounting.analysed = {'M1': 'M1.vcf', 'C1': 'C1.vcf'}
    accounting.remove_unmatched = ['NotHere1', 'NotHere2']
    summary.accounting = accounting
    summary.args = SimpleNamespace(remove_by_name='/x/remove_from_analysis.xlsx')
    summary.validation_results = None
    summary.metadata_problem = summary.reference_note = None
    summary.name_collisions, summary.metadata_ambiguous = {}, {}
    summary.groupings_dict = {'G1': {'M1': None, 'C1': None}, 'G2': {'M1': None}}
    summary.warnings()
    summary.sample_accounting()
    summary.groupings()
    html = '\n'.join(summary.out)
    check('no remove-list box, mixed-sample section or multi-base section',
          'matched no VCF' not in html and 'NotHere1' not in html
          and 'Possible mixed samples' not in html and 'multi-base' not in html, html[:300])
    m1 = next((line for line in summary.out if '>M1<' in line), '')
    c1 = next((line for line in summary.out if '>C1<' in line), '')
    check('a mixed sample is highlighted in the Groupings table, its counts after its groups',
          "<span class='mixed'>M1</span>" in m1
          and m1.index('G2') < m1.index('58 confident mixed calls against 5 passing SNPs'), m1)
    check('a clonal sample is not highlighted', c1 and 'mixed' not in c1, c1)
    check('the Groupings table explains the highlight once',
          html.count("<span class='mixed'>Highlighted</span>") == 1)


def section_truncated(directory):
    print('\n[15] a truncated VCF is excluded (D-06)')
    from vsnp3_input_validator import scan_vcf_header, classify_scan, InputValidator
    full = HEADER + snp('NC_1', 100, 'A', 'G') + snp('NC_1', 200, 'C', 'T') + snp('NC_2', 50, 'G', 'A')
    cut_mid = write(directory, 'cutmid_zc.vcf', full[:-25])
    cut_line = write(directory, 'cutline_zc.vcf', full[:full.rindex('NC_2\t50')])
    whole = write(directory, 'whole_zc.vcf', full)
    category, message = classify_scan(scan_vcf_header(cut_mid))
    check('a file cut part-way through a record is corrupted: truncated',
          category == 'corrupted' and 'truncated' in message, f'{category}: {message}')
    check('a complete file is valid', classify_scan(scan_vcf_header(whole))[0] == 'valid')
    check('a cut exactly on a line boundary cannot be seen (documented limit)',
          classify_scan(scan_vcf_header(cut_line))[0] == 'valid')
    results = InputValidator().validate_vcf_list([whole, cut_mid])
    check('validation excludes it with the reason',
          cut_mid not in results.valid_files and 'truncated' in results.reason(cut_mid)
          and whole in results.valid_files, str(results.reasons))


def section_step1_report():
    print('\n[7] usability verdicts stay out of the workbook')
    try:
        import pandas as pd
        from vsnp3_file_setup import Excel_Stats, fastq_usability_verdict, reference_usability_verdict, HtmlReport
    except ImportError as e:
        print(f'  SKIP  ({e})')
        return
    check('verdict names the measurement',
          fastq_usability_verdict(35.0, 0.05, 96.0, 88.0) == ('Questionable', 'average depth 35.0x < 40.0x'))
    check('acceptable when all limits met',
          reference_usability_verdict(0.995, 0.5, 0.01)[0] == 'Acceptable')
    check('string inputs tolerated', fastq_usability_verdict('120', '0.05', '96', '88')[0] == 'Acceptable')
    cwd = os.getcwd()
    with tempfile.TemporaryDirectory() as tmp:
        try:
            os.chdir(tmp)
            xs = Excel_Stats('T-1')
            xs.excel_dict['Average Depth'] = '88.2'
            xs.excel_dict['FASTQ Usability'] = 'Questionable (average depth 35.0x < 40.0x)'
            xs.excel_dict['Reference Usability'] = 'Acceptable (all limits met)'
            try:
                xs.post_excel()
            except ImportError as e:
                print(f'  SKIP  workbook ({e})')
                return
            df = pd.read_excel(xs.excel_filename)
            check('workbook has no usability columns',
                  not any('Usability' in str(c) for c in df.columns), str(list(df.columns)))
            check('workbook keeps the measurements', 'Average Depth' in df.columns)
            report = HtmlReport('T-1')
            try:
                path = report.write(dict(xs.excel_dict))
            except ImportError as e:
                print(f'  SKIP  report ({e})')
                return
            body = open(path, encoding='utf-8').read()
            check('report shows the verdict and its reason',
                  'Questionable' in body and 'average depth 35.0x' in body)
        finally:
            os.chdir(cwd)


def section_longread(directory):
    print('\n[16] long-read (nanopore) evidence QUAL (vsnp3_longread.py)')
    import math
    import pandas as pd
    from vsnp3_longread import (evidence_qual, per_read_weight, homopolymer_touching,
                                one_strand_only, apply_model, model_header_lines, mixed_by_evidence)
    from vsnp3_version import (QUAL_THRESHOLD, N_THRESHOLD, LONGREAD_VARIANT_PRIOR,
                               LONGREAD_HOMOPOLYMER_MIN, LONGREAD_STRAND_MIN_ALT,
                               LONGREAD_CALLER_MIN_QUAL)
    e = 0.015
    w = per_read_weight(e)
    check('30 clean ALT reads at 1.5% error are a confident call (QUAL > 150)',
          evidence_qual(0, 0, 15, 15, e) > QUAL_THRESHOLD, str(evidence_qual(0, 0, 15, 15, e)))
    check('QUAL grows with every supporting read',
          all(evidence_qual(0, 0, k, k, e) < evidence_qual(0, 0, k + 1, k + 1, e) for k in range(1, 40)))
    check('QUAL falls with every reference read',
          all(evidence_qual(k, 0, 10, 10, e) > evidence_qual(k + 1, 0, 10, 10, e) for k in range(0, 12)))
    check('a 15/15 mixed site is a confident (IUPAC) call',
          evidence_qual(8, 7, 8, 7, e) > QUAL_THRESHOLD, str(evidence_qual(8, 7, 8, 7, e)))
    check('3 ALT reads among 30 fall below the N band (reference in step 2)',
          evidence_qual(14, 13, 2, 1, e) < N_THRESHOLD, str(evidence_qual(14, 13, 2, 1, e)))
    check('two ALT reads alone never pass', evidence_qual(0, 0, 1, 1, e) < N_THRESHOLD)
    check('no reads scores 0', evidence_qual(0, 0, 0, 0, e) == 0.0)
    closed = 30 * w + 10 * math.log10(LONGREAD_VARIANT_PRIOR / 2) - 10 * math.log10(1 - LONGREAD_VARIANT_PRIOR)
    check('a clean site scores the per-read weight times reads, less the prior (freebayes scale)',
          abs(evidence_qual(0, 0, 15, 15, e) - closed) < 0.5,
          f'{evidence_qual(0, 0, 15, 15, e)} vs {closed:.1f}')
    check('an Illumina-like error rate weighs a read close to freebayes (25-30)',
          25 <= per_read_weight(0.003) <= 30, str(per_read_weight(0.003)))
    check('a nanopore read is worth less than an Illumina read', per_read_weight(0.03) < per_read_weight(0.003))
    check('deep sites do not overflow', evidence_qual(0, 0, 5000, 5000, e) > 10000)
    check('four reference-looking reads among thirty are homozygous at a 2.3% error rate',
          not mixed_by_evidence(3, 1, 15, 11, 0.023))
    check('a 60/40 site is mixed', mixed_by_evidence(6, 6, 9, 9, 0.023))
    check('a half-and-half site is mixed', mixed_by_evidence(8, 7, 8, 7, 0.023))
    check('a clean site is homozygous', not mixed_by_evidence(0, 0, 15, 15, 0.023))
    check('no reads is not mixed', not mixed_by_evidence(0, 0, 0, 0, 0.023))

    seq = 'ACGTAAAAAACGT'                       # a run of six A at index 4..9
    check('a call inside a homopolymer run sees its length', homopolymer_touching(seq, 6) == 6)
    check('a call on the base before the run counts as beside it', homopolymer_touching(seq, 3) == 6)
    check('a call on the base after the run counts as beside it', homopolymer_touching(seq, 10) == 6)
    check('a call away from any run sees 1',
          homopolymer_touching(seq, 1) == 1 and homopolymer_touching(seq, 12) == 1)
    check('out of range is 0', homopolymer_touching('', 0) == 0 and homopolymer_touching(seq, 99) == 0)

    check('ALT on the forward strand only, both strands sequenced: one-strand',
          one_strand_only(5, 5, 6, 0) and one_strand_only(3, 3, 0, 4))
    check('ALT on both strands: not one-strand', not one_strand_only(5, 5, 3, 3))
    check('reverse strand not sequenced at all: the rule does not apply', not one_strand_only(10, 0, 6, 0))
    check(f'fewer than {LONGREAD_STRAND_MIN_ALT} ALT reads: the rule does not apply',
          not one_strand_only(5, 5, 2, 0))

    # A bcftools-style VCF through apply_model, then through every reader.
    fasta = write(directory, 'lr_ref.fasta',
                  '>NC_1 test\n' + ('ACGT' * 25) + 'GCAAAAAAGC' + ('TGCA' * 25) + '\n'
                  '>NC_2 test\n' + ('CGTA' * 30) + '\n')
    # NC_1: positions 1-100 ACGT..., 101 G, 102 C, 103-108 AAAAAA, 109 G, 110 C, 111.. TGCA...
    hdr = ('##fileformat=VCFv4.2\n##reference=ref.fasta\n'
           '##contig=<ID=NC_1,length=210>\n##contig=<ID=NC_2,length=120>\n'
           '##INFO=<ID=DP,Number=1,Type=Integer,Description="depth">\n'
           '##INFO=<ID=DP4,Number=4,Type=Integer,Description="ref-fwd, ref-rev, alt-fwd, alt-rev">\n'
           '##INFO=<ID=MQ,Number=1,Type=Integer,Description="RMS mapping quality">\n'
           '##INFO=<ID=AC,Number=A,Type=Integer,Description="allele count">\n'
           '#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tsample\n')

    def rec(pos, ref, alt, qual, dp4, ac, mq=60, gt=None):
        gt = gt or ('1/1' if ac == '2' else '0/1')
        dp = sum(int(x) for x in dp4.split(','))
        return (f'NC_1\t{pos}\t.\t{ref}\t{alt}\t{qual}\t.\tDP={dp};DP4={dp4};MQ={mq};AC={ac};AN=2'
                f'\tGT:PL:AD\t{gt}:255,90,0:{dp4.split(",")[0]},{dp4.split(",")[2]}\n')
    raw = (hdr
           + rec(50, 'C', 'T', 225.4, '0,0,15,15', '2')          # clean, confident
           + rec(60, 'T', 'A', 228.0, '1,2,20,0', '2')           # ALT forward only: N band
           + rec(70, 'C', 'G', 15.0, '0,0,10,10', '2')           # caller QUAL below the gate: dropped
           + 'NC_1\t80\t.\tGT\tG\t200\t.\tDP=30;DP4=0,0,15,15;MQ=60;AC=2;AN=2;INDEL\tGT\t1/1\n'  # indel: dropped
           + rec(104, 'A', 'G', 220.0, '0,0,15,15', '2')         # inside the AAAAAA run: N band
           + rec(102, 'C', 'T', 220.0, '0,0,15,15', '2')         # the base before the run: N band
           + rec(106, 'A', 'C', 25.0, '10,10,3,2', '1')          # weak, inside the run: raised to N
           + rec(40, 'G', 'C', 195.4, '3,1,15,11', '1')          # bcftools het; the model says homozygous
           + rec(44, 'G', 'C', 200.0, '0,0,14,14', '1', gt='0/1')  # bcftools het on a clean site: homozygous
           + rec(90, 'C', 'A', 100.0, '7,8,8,7', '1')            # mixed, confident
           + rec(94, 'C', 'A', 30.0, '10,10,2,1', '1')           # mixed, thin: below N band
           + rec(98, 'C', 'A,T', 150.0, '0,0,10,10', '1,1')      # multi-allelic single bases: kept
           + 'NC_1\t20\t.\tA\tG\t200\t.\tDP=20;MQ=60;AC=2;AN=2\tGT:AD\t1/1:0,20\n')  # no DP4: AD used
    raw_path = write(directory, 'lr_raw.vcf', raw)
    out_path = os.path.join(directory, 'lr-ont_filtered_hapall_nanopore.vcf')
    before = sha(raw_path)
    counts = apply_model(raw_path, out_path, fasta, e, 'test', 'bcftools test')
    check('the caller VCF is not modified', sha(raw_path) == before)
    check('records are counted: 13 in, 1 indel, 1 below the caller gate, 11 written',
          (counts['records'], counts['not_snp'], counts['below_caller_qual'], counts['written'],
           counts['no_depth']) == (13, 1, 1, 11, 0), str(counts))
    check('two genotypes decided differently from bcftools', counts['genotype_changed'] == 2, str(counts))
    check('one-strand and homopolymer flags counted, every flagged record placed in the N band',
          (counts['flag_strand'], counts['flag_homopolymer'], counts['limited']) == (1, 3, 4), str(counts))
    check('thin calls are counted as raised to the floor', counts['floored'] == 2, str(counts))
    text = open(out_path).read()
    lines = text.splitlines()
    headers = [l for l in lines if l.startswith('##')]
    check('header names the platform and the model',
          '##vsnp3_platform=nanopore' in headers
          and any(l.startswith('##vsnp3_longread_model=error_rate=0.0150;') for l in headers)
          and any(l.startswith('##vsnp3_longread_caller=bcftools test') for l in headers)
          and any('ID=VSNP_RAWQ' in l for l in headers) and any('ID=VSNP_FLAG' in l for l in headers))
    check('header lines come before the column line',
          lines.index('##vsnp3_platform=nanopore') < next(i for i, l in enumerate(lines) if l.startswith('#CHROM')))
    body_rows = {int(l.split('\t')[1]): l.split('\t') for l in lines if not l.startswith('#')}
    r50 = body_rows[50]
    check('a clean call carries the evidence QUAL and keeps the caller QUAL in INFO',
          float(r50[5]) == evidence_qual(0, 0, 15, 15, e) and 'VSNP_RAWQ=225.4' in r50[7]
          and r50[6] == '.' and 'VSNP_FLAG' not in r50[7], str(r50))
    r60 = body_rows[60]
    check('a strong one-strand call is held at QUAL 149 with the reason in FILTER and INFO',
          float(r60[5]) == QUAL_THRESHOLD - 1 and r60[6] == 'vsnp_strand'
          and 'VSNP_FLAG=strand' in r60[7] and float(r60[7].split('VSNP_EVQ=')[1].split(';')[0]) > QUAL_THRESHOLD,
          str(r60))
    check('a call inside and a call beside a homopolymer run are held at 149',
          float(body_rows[104][5]) == QUAL_THRESHOLD - 1 and 'homopolymer' in body_rows[104][6]
          and float(body_rows[102][5]) == QUAL_THRESHOLD - 1 and 'homopolymer' in body_rows[102][6])
    check('a weak flagged call is raised to the N band, not left as a reference call',
          float(body_rows[106][5]) == N_THRESHOLD and 'homopolymer' in body_rows[106][6]
          and float(body_rows[106][7].split('VSNP_EVQ=')[1].split(';')[0]) < N_THRESHOLD, str(body_rows[106]))
    check('a run shorter than the limit is not flagged',
          f'homopolymer_min={LONGREAD_HOMOPOLYMER_MIN}' in text and body_rows[50][6] == '.')
    check('dropped records are absent', 70 not in body_rows and 80 not in body_rows)
    check('a confident mixed call keeps AC=1 and scores above 150',
          float(body_rows[90][5]) > QUAL_THRESHOLD and 'AC=1' in body_rows[90][7])
    check('a thin mixed call is raised to the N floor, with its evidence kept in INFO',
          float(body_rows[94][5]) == N_THRESHOLD
          and float(body_rows[94][7].split('VSNP_EVQ=')[1].split(';')[0]) < N_THRESHOLD, str(body_rows[94]))
    check('without DP4 the AD field is used', 20 in body_rows and float(body_rows[20][5]) > QUAL_THRESHOLD)
    check('the genotype and FORMAT columns are untouched where the model agrees',
          r50[8] == 'GT:PL:AD' and r50[9].startswith('1/1:') and body_rows[90][9].startswith('0/1:')
          and 'VSNP_CALLER_GT' not in r50[7] and 'VSNP_CALLER_GT' not in body_rows[90][7])
    r40 = body_rows[40]
    check('a clonal site bcftools called mixed is homozygous: AC=2, GT 1/1, the caller genotype kept',
          'AC=2' in r40[7] and r40[9].startswith('1/1:') and 'VSNP_CALLER_GT=0/1' in r40[7]
          and r40[9].endswith(':3,15') and float(r40[5]) > QUAL_THRESHOLD, str(r40))
    check('the AC rewrite leaves the rest of INFO as it was',
          r40[7].startswith('DP=30;DP4=3,1,15,11;MQ=60;AC=2;AN=2;VSNP_RAWQ='), str(r40[7]))
    check('a clean site bcftools called mixed is homozygous too',
          'AC=2' in body_rows[44][7] and body_rows[44][9].startswith('1/1:'))
    check('a multi-allelic record keeps its genotype', 'AC=1,1' in body_rows[98][7])

    # Every reader takes the file as it takes a freebayes VCF.
    from vsnp3_step2 import parse_vcf
    from vsnp3_group_on_defining_snps import Group
    from vsnp3_zero_coverage import step2_snp_counts, confident_mixed_count
    from vsnp3_group_reporter import GroupReporter
    _, df, stats, err = parse_vcf(out_path)
    check('step 2 reads the long-read VCF', err is None and df is not None, str(err))
    if df is not None:
        check('step 2 records the platform and model from the header',
              stats['platform'] == 'nanopore' and str(stats['longread_model']).startswith('error_rate=0.0150'),
              str({k: stats[k] for k in ('platform', 'longread_model')}))
        row = df[df['abs_pos'] == 'NC_1:50'].iloc[0]
        check('QUAL, AC, DP and MQ are read; FILTER is not a column',
              row['QUAL'] == int(evidence_qual(0, 0, 15, 15, e)) and row['AC'] == 2 and row['DP'] == 30
              and row['MQ'] == 60 and 'FILTER' not in df.columns, str(row.to_dict()))
        group = Group.__new__(Group)
        group.qual_threshold, group.mq_threshold = 150, 56
        passing = group._passing(df)
        vcf_df = pd.read_csv(out_path, sep='\t', header=None, comment='#',
                             names=['CHROM', 'POS', 'ID', 'REF', 'ALT', 'QUAL', 'FILTER', 'INFO', 'FORMAT', 'Sample'])
        thin = df[df['abs_pos'] == 'NC_1:94'].iloc[0]
        check('step 2 reads the floored call as N, not as the reference',
              N_THRESHOLD <= thin['QUAL'] < QUAL_THRESHOLD, str(thin.to_dict()))
        check('the clean AC=2 call and the two re-genotyped calls pass; the limited calls do not',
              sorted(passing['abs_pos']) == ['NC_1:20', 'NC_1:40', 'NC_1:44', 'NC_1:50'],
              passing['abs_pos'].tolist())
        check('step 1 counts what step 2 takes, on a long-read VCF too',
              step2_snp_counts(vcf_df)[1] == len(passing) and confident_mixed_count(vcf_df) == group._confident_mixed(df),
              f'{step2_snp_counts(vcf_df)} {len(passing)} {confident_mixed_count(vcf_df)} {group._confident_mixed(df)}')
        limited = df[df['abs_pos'].isin(['NC_1:60', 'NC_1:102', 'NC_1:104', 'NC_1:106'])]
        check('limited calls sit in the N band for make_groupings',
              ((limited['QUAL'] >= N_THRESHOLD) & (limited['QUAL'] < QUAL_THRESHOLD)).all())
    gr = GroupReporter.read_vcf(None, out_path)
    check('the group reporter reads it', len(gr) == 11 and int(gr.loc[gr['POS'] == 50, 'QUAL'].iloc[0]) > 150
          and int(gr.loc[gr['POS'] == 60, 'AC'].iloc[0]) == 2, str(gr.head()))
    hl = model_header_lines('bcftools x', 0.05, 'default')
    check('model header lines carry every parameter',
          any('weight_per_read=12.8' in l for l in hl) and any(f'prior={LONGREAD_VARIANT_PRIOR}' in l for l in hl)
          and any(f'floor_qual={N_THRESHOLD}' in l for l in hl) and any('genotype=evidence model' in l for l in hl)
          and any(f'<= {LONGREAD_CALLER_MIN_QUAL}' in l for l in hl), str(hl))



def section_robustness(directory):
    print('\n[17] robustness fixes of 2026-10-16 (audit items H1-H5, M1-M8, L6)')
    import numpy as np
    import pandas as pd
    import vsnp3_html_tree as vht
    from vsnp3_fasta_to_snps_table import Tables
    from vsnp3_group_on_defining_snps import Group, DefiningSnpsError, NEWICK_UNSAFE
    from vsnp3_file_setup import worker_count

    # H1: leaf order from a tree whose names carry <digit>.<digit>
    tree = os.path.join(directory, 'leaf_order.tre')
    with open(tree, 'w') as f:
        f.write("((22-12345.1_zc.vcf:0.001,A_zc.vcf:0.002):0.1,"
                "(NZ_CP0123.1.vcf:0.003,'B C_zc.vcf':0.004):0.2,root:0.0);\n")
    order = Tables.leaf_order(tree, 'T')
    check('leaf order keeps names containing digit.digit, in file order, root first',
          order == ['root', '22-12345.1_zc.vcf', 'A_zc.vcf', 'NZ_CP0123.1.vcf', 'B C_zc.vcf'], str(order))
    with open(tree, 'w') as f:
        f.write("")             # an empty tree file (Bio.Phylo reads most text as a one-leaf tree)
    try:
        Tables.leaf_order(tree, 'T')
        check('an unreadable tree is an error naming the group', False)
    except RuntimeError as e:
        check('an unreadable tree is an error naming the group', '[T]' in str(e), str(e))

    # H2: the 'no annotations' row is not a sample
    frame = _table_frame(['root', 'S1', 'S2', 'MQ', 'no annotations']).reset_index()
    frame.columns = ['Sample'] + list(frame.columns[1:])
    matrix = vht.SnpMatrix(vht.select_position_columns(frame))
    check('"no annotations" row is not a leaf of the HTML tree',
          matrix.sample_names == ['S1', 'S2'], str(matrix.sample_names))

    # M5: count_leaves on a chain deeper than the C-level recursion limit
    top = node = vht.Node('n0')
    for i in range(1, 6000):
        child = vht.Node(f'n{i}')
        child.parent = node
        node.children.append(child)
        node = child
    try:
        check('count_leaves survives a 6,000-deep chain', vht.count_leaves(top) == 1)
    except RecursionError:
        check('count_leaves survives a 6,000-deep chain', False, 'RecursionError')

    # H3: blank metadata cells
    xlsx = os.path.join(directory, 'meta_blank.xlsx')
    pd.DataFrame([['a_zc.vcf', 'A name'], ['b_zc.vcf', None], [None, 'orphan'],
                  ['c_zc.vcf', 'C:x']]).to_excel(xlsx, header=False, index=False)
    meta = Group.read_metadata(xlsx)
    check('rows with a blank cell are skipped and counted',
          meta.attrs.get('blank_rows') == 2 and meta['file_name'].tolist() == ['a_zc.vcf', 'c_zc.vcf'],
          f"{meta.attrs} {meta['file_name'].tolist()}")
    check('no sample is named "nan"', 'nan' not in meta['metadata'].tolist(), str(meta['metadata'].tolist()))
    name, key, hits = Group.resolve_sample_name_detail('b_zc.vcf', meta, True)
    check('a file whose row was blank keeps its file-derived name', key is None and name == 'b', f'{name} {key}')

    # H4: fallback names RAxML can carry
    check('safe_fallback_name leaves a plain name alone', Group.safe_fallback_name('S1_zc.vcf') == 'S1_zc.vcf')
    check('safe_fallback_name replaces Newick-unsafe characters',
          Group.safe_fallback_name("a b(1):x,y;'z'") == 'a_b_1_x_y_z', Group.safe_fallback_name("a b(1):x,y;'z'"))
    check('a dot is not replaced (RAxML accepts it; metadata names already lose it)',
          Group.safe_fallback_name('NZ_CP0123.1') == 'NZ_CP0123.1' and not NEWICK_UNSAFE.search('.'))

    # M1: worksheet checks
    g = Group.__new__(Group)
    g.worksheet_notes = []
    df_ok = pd.DataFrame({'all': ['x', 'NC_1:5'], 'NC_1:100': ['G1', 'NC_1:7'], 'NC_1:200': ['G1', None],
                          'NC_1:300': ['G2', 12345]})
    out = g._check_defining_snps('w.xlsx', df_ok, {'NC_1:100': 'G1', 'NC_1:200': 'G1', 'NC_1:300': 'G2'},
                                 {'G1': ['NC_1:7'], 'G1.1': ['NC_1:8'], 'G2': [12345]}, ['NC_1:5'])
    check('duplicate group names are merged, including their filter columns',
          out.get('G1') == ['NC_1:7', 'NC_1:8'] and 'G1.1' not in out, str(out))
    check('the merge and the ignored filter entry are noted for the summary',
          len(g.worksheet_notes) == 2 and 'merged' in g.worksheet_notes[0] and '12345' in g.worksheet_notes[1],
          str(g.worksheet_notes))
    for bad_df, bad_dict, what in [
            (pd.DataFrame({'all': ['x'], 12345: ['G1']}), {12345: 'G1'}, 'an integer position header'),
            (pd.DataFrame({'all': ['x'], 'NC_1:100': [float('nan')]}), {'NC_1:100': float('nan')}, 'a blank group name')]:
        try:
            Group.__new__(Group)._check_defining_snps('w.xlsx', bad_df, bad_dict, {}, [])
            check(f'{what} is a DefiningSnpsError', False)
        except DefiningSnpsError as e:
            check(f'{what} is a DefiningSnpsError', 'w.xlsx' in str(e), str(e))
        except Exception as e:                                  # noqa: BLE001
            check(f'{what} is a DefiningSnpsError', False, f'{type(e).__name__}: {e}')

    # H5 + H4: parsimony test and a sample with no determined base
    def frame_of(calls):
        rows = []
        for pos, (ref, alt, qual, ac) in calls.items():
            rows.append(('NC_1', pos, ref, alt, qual, ac, 60 if qual is not None else None))
        df = pd.DataFrame(rows, columns=['CHROM', 'POS', 'REF', 'ALT', 'QUAL', 'AC', 'MQ'])
        df['QUAL'] = pd.array(df['QUAL'], dtype='Int64')
        df['MQ'] = pd.array(df['MQ'], dtype='Int64')
        df['AC'] = df['AC'].astype(np.int64)
        df['DP'] = 50
        df['CIGAR'] = np.nan
        df['abs_pos'] = df['CHROM'] + ':' + df['POS'].astype(str)
        return df
    g = Group.__new__(Group)
    g.debug = False
    g.n_threshold, g.qual_threshold, g.mq_threshold = 50, 150, 56
    g.ambigious_lookup = {'AG': 'R', 'GA': 'R', 'CT': 'Y', 'TC': 'Y'}
    g.sample_stats = {}
    g.embedded_cells, g.embedded_now_uninformative, g._embedded_cache = {}, {}, {}
    from collections import defaultdict
    g.embedded_cells, g.embedded_now_uninformative = defaultdict(int), defaultdict(set)
    g.dataframes_names_updated = {
        's1': frame_of({100: ('A', 'G', 900, 2), 200: ('C', 'T', 900, 2), 300: ('G', 'A', 900, 2)}),
        's2': frame_of({100: ('A', 'G', 900, 2), 200: ('C', 'C', 900, 2), 300: ('G', 'A', 900, 2)}),
        's3': frame_of({100: ('A', 'G', 900, 2), 200: ('N', '.', None, 0), 300: ('G', 'A', 900, 2)}),
        'blank': frame_of({100: ('N', '.', None, 0), 200: ('N', '.', None, 0), 300: ('N', '.', None, 0)}),
    }
    ess = {k: g._passing(v) for k, v in g.dataframes_names_updated.items()}
    result, name = g.make_groupings(('G', ess))
    seqs = {k: ''.join(v['ALT'].tolist()) for k, v in result.items()}
    positions = list(result['root']['abs_pos']) if 'root' in result else []
    check('positions where every sample carries the same base leave the table',
          positions == ['NC_1:200'], f'{positions} {seqs}')
    check('samples with no determined base at the informative positions are left out and recorded',
          'blank' not in result and 's3' not in result
          and g.undetermined_in_group.get('G') == ['blank', 's3'],
          f'{sorted(result)} {dict(g.undetermined_in_group)}')
    check('the other samples align with the root', set(seqs) == {'s1', 's2', 'root'}
          and seqs['s1'] == 'T' and seqs['s2'] == 'C' and seqs['root'] == 'C', str(seqs))

    # M2: the reader's single_base column equals the fallback
    df = g.dataframes_names_updated['s1'].copy()
    fallback = Group._single_base(df)
    df['single_base'] = (df['REF'].str.len() == 1) & (df['ALT'].str.len() == 1)
    check('_single_base uses the reader column when present and agrees with the fallback',
          Group._single_base(df).tolist() == fallback.tolist())

    # M4: HTML-tree timeout
    class _T:
        pass
    t = _T()
    t.table_to_tree_df = pd.DataFrame(index=range(400))
    saved = os.environ.pop('VSNP3_HTML_TREE_TIMEOUT', None)
    try:
        check('timeout is 900 s for a small group', Group._html_tree_timeout(t) == 900)
        t.table_to_tree_df = pd.DataFrame(index=range(10000))
        check('timeout grows with the group (10,000 samples: about 6,200 s)',
              6000 < Group._html_tree_timeout(t) < 6400, str(Group._html_tree_timeout(t)))
        os.environ['VSNP3_HTML_TREE_TIMEOUT'] = '7000'
        check('VSNP3_HTML_TREE_TIMEOUT raises the limit', Group._html_tree_timeout(t) == 7000)
        os.environ['VSNP3_HTML_TREE_TIMEOUT'] = 'junk'
        check('a junk VSNP3_HTML_TREE_TIMEOUT falls back to the default', Group._html_tree_timeout(t) > 6000)
    finally:
        os.environ.pop('VSNP3_HTML_TREE_TIMEOUT', None)
        if saved is not None:
            os.environ['VSNP3_HTML_TREE_TIMEOUT'] = saved

    # worker_count (item 50) and the remove list's entries (M6)
    saved = os.environ.pop('VSNP3_MAX_CPUS', None)
    try:
        check('worker_count is at least 1 and at most 32 by default', 1 <= worker_count() <= 32)
        check('worker_count never exceeds the task count', worker_count(2) <= 2)
        os.environ['VSNP3_MAX_CPUS'] = '1'
        check('VSNP3_MAX_CPUS caps worker_count', worker_count() == 1)
    finally:
        os.environ.pop('VSNP3_MAX_CPUS', None)
        if saved is not None:
            os.environ['VSNP3_MAX_CPUS'] = saved
    from vsnp3_remove_from_analysis import Remove_From_Analysis
    rm_xlsx = os.path.join(directory, 'remove.xlsx')
    pd.DataFrame([['present_zc'], ['missing_sample']]).to_excel(rm_xlsx, header=False, index=False)
    rfa = Remove_From_Analysis(working_directory=directory, excel_remove=rm_xlsx, extension='vcf')
    check('the remove list keeps its worksheet entries for reporting',
          rfa.entries == ['present_zc', 'missing_sample'], str(getattr(rfa, 'entries', None)))

    # M8: Excel merge tool strips the basename only
    import vsnp3_excel_merge_defining_snps as merge
    src = merge.Expand_Range.ex_range.__code__
    check('the merge tool derives its output name from the basename',
          'splitext' in open(merge.__file__).read())

    # L6: a name starting with '=' is written as text
    import openpyxl
    tables = _tables()
    table = tables.json_round_trip(_table_frame(['root', '=S1', 'S2', 'MQ', 'annotation']))
    path = os.path.join(directory, 'formula_name.xlsx')
    tables.excel_formatter(table, path)
    ws = openpyxl.load_workbook(path).active
    cell = next((c for row in ws.iter_rows() for c in row if c.value == '=S1'), None)
    check('a sample name starting with "=" is a text cell, not a formula',
          cell is not None and cell.data_type == 's', f'{cell.data_type if cell else None}')


def main():
    print('=' * 78)
    print('vSNP3 step 2 input tests')
    print('=' * 78)
    with tempfile.TemporaryDirectory() as tmp:
        paths = make_set(tmp)
        section_validator(tmp, paths)
        section_reader(tmp, paths)
        section_fix_vcfs(tmp, paths)
        section_group(tmp, paths)
        section_reference_split(tmp)
        section_workbook_format(tmp)
        section_na_names(tmp)
        section_html_tree_failures(tmp)
        section_step1_notice(tmp)
        section_multibase(tmp)
        section_mixed_signature(tmp)
        section_truncated(tmp)
        section_longread(tmp)
        section_robustness(tmp)
    section_summary_layout()
    section_metadata()
    section_timeout()
    section_step1_report()
    print('\n' + '=' * 78)
    print(f'{len(PASS)} passed, {len(FAIL)} failed')
    if FAIL:
        print('FAILED:')
        for name in FAIL:
            print(f'  - {name}')
    print('=' * 78)
    return 1 if FAIL else 0


if __name__ == '__main__':
    import multiprocessing
    multiprocessing.set_start_method('spawn', True)
    sys.exit(main())
